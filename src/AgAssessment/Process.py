# --------------------------------------------------------------------------------
# Name:        Process
# Purpose:     This tool performs all of the ag assessment calculations
#
# License:     Contextual Copyleft AI (CCAI) License v1.0.
#              Full license in LICENSE file.
# --------------------------------------------------------------------------------

import json
import pathlib

import openpyxl

import arcpy
from helpers.logging import error, log, warn
from helpers.parameter import sanitize, set_required_parameter
from helpers.parameter import validate_spatial_reference as validate
from helpers.tool import (
    license,
    reload_module,
)
from helpers.tool import setup_environment as setup

from .DefineParcels import AG_ASSESSMENT_GDB_NAME


class Process:
    def __init__(self):
        """Define the tool (tool name is the name of the class)."""
        self.label = "3. Process"
        self.category = "Automated Ag Assessment"
        self.description = "Run after splitting parcels into use areas"

    def getParameterInfo(self):
        """Define parameter definitions"""
        param0 = arcpy.Parameter(
            displayName="Soils",
            name="soils",
            datatype="GPFeatureLayer",
            parameterType="Required",
            direction="Input")
        param0.filter.list = ["Polygon"]

        param1 = arcpy.Parameter(
            displayName="Soils MUSYM Field",
            name="soils_musym_field",
            datatype="GPString",
            parameterType="Required",
            direction="Input",
        )
        param1.filter.type = "ValueList"
        param1.filter.list = []

        param2 = arcpy.Parameter(
            displayName="Soils MUKEY Field",
            name="soils_mukey_field",
            datatype="GPString",
            parameterType="Required",
            direction="Input",
        )
        param2.filter.type = "ValueList"
        param2.filter.list = []

        params = [param0, param1, param2]
        return params

    def set_dependent_layers(self, parameters):
        fields = [f.name for f in arcpy.ListFields(parameters[0].value)]
        parameters[1].enabled = True
        parameters[2].enabled = True
        parameters[1].filter.list = fields
        parameters[2].filter.list = fields
        if "MUSYM" in fields:
            parameters[1].value = "MUSYM"
        if "MUKEY" in fields:
            parameters[2].value = "MUKEY"

    def updateParameters(self, parameters):
        # get soils MUSYM nad MUKEY field
        if not parameters[0].hasBeenValidated:
            if parameters[0].value:
                self.set_dependent_layers(parameters)
            else:
                parameters[1].enabled = False
                parameters[2].enabled = False

                # check if we have a default layer
                project = arcpy.mp.ArcGISProject("Current")
                active_map = project.activeMap
                lyrs = active_map.listLayers()
                for lyr in lyrs:
                    if lyr.isGroupLayer:
                        continue
                    elif "soil" in lyr.name.lower():
                        parameters[0].value = lyr.longName
                        self.set_dependent_layers(parameters)
                        break

    def updateMessages(self, parameters):
        """Modify the messages created by internal validation for each tool parameter."""
        # make newly toggled on parameters required
        set_required_parameter(parameters[0].value, parameters[1])
        set_required_parameter(parameters[0].value, parameters[2])

        validate(parameters)

    def isLicensed(self):
        """Set whether the tool is licensed to execute."""
        return license()

    @reload_module(__name__)
    def execute(self, parameters, messages):
        """The source code of the tool."""
        # Setup
        log("setting up project")
        project, active_map = setup()
        project_dir = project.homeFolder
        cache_file_path = f"{project_dir}/.ag_cache.json"

        # read in json
        log("reading in cache")
        cache = {}
        with open(cache_file_path) as file:
            cache = json.load(file)
        parcels = cache["parcels"]
        output_folder = cache["output_folder"]

        # Parameters
        log("reading in parameters")
        soil_layer = parameters[0].value
        soils_musym = parameters[1].value
        soils_mukey = parameters[2].value

        # check for geodatabase and set it as workspace
        db_path = f"{project.homeFolder}\\{AG_ASSESSMENT_GDB_NAME}.gdb"
        if not arcpy.Exists(db_path):
            error(f"Ag assessment geodatase {db_path} does not exist. Please start over with step 1.")
        arcpy.env.workspace = db_path

        # collect layouts to be able to close and redisplay later
        layouts = []
        log("iterating through parcels and processing")
        for parcel in parcels:
            # find map of parcel
            m = None
            try:
                m = project.listMaps(parcel)[0]
            except Exception:
                warn(f"unable to find map for {parcel}, results may be incomplete")
                continue

            # Clear selection
            m.clearSelection()

            # find layout
            lyt = None
            try:
                lyt = project.listLayouts(parcel)[0]
                layouts.append(lyt)
            except Exception:
                warn(f"couldn't find layout for parcel {parcel}, results may be incomplete")
                continue

            # Helper variables
            soils_layers = []
            use_layers = []
            tables = []

            # Start work
            log(f"processing {parcel}")
            lyrs = m.listLayers()
            lyr_types = set()
            for lyr in lyrs:
                # Update symbology
                lyr_type = ""

                # find layer types
                if "Agland" == lyr.name:
                    use_layers.append(lyr)
                    lyr_type = "Agland"
                elif "NonAg" == lyr.name:
                    use_layers.append(lyr)
                    lyr_type = "NonAg"
                elif "Forest" == lyr.name:
                    use_layers.append(lyr)
                    lyr_type = "Forest"
                else:
                    continue
                lyr_types.add(lyr_type)

                # Create clip layer
                new_layer_name = f"{lyr_type}_{parcel}"
                new_layer_path = f"{arcpy.env.workspace}\\{sanitize(new_layer_name)}_soils"
                arcpy.analysis.Clip(soil_layer, lyr, new_layer_path)

                # Dissolve duplicate MUSYMs
                dissolve_layer_path = f"{arcpy.env.workspace}\\{sanitize(new_layer_name)}_soils_dissolved"
                arcpy.management.Dissolve(new_layer_path, dissolve_layer_path, [soils_musym, soils_mukey])

                # Add to map
                new_layer = m.addDataFromPath(dissolve_layer_path)
                soils_layers.append(new_layer)
                new_layer.name = new_layer_name

                # Add acreage field
                if "Acres" not in [f.name for f in arcpy.ListFields(new_layer.dataSource)]:
                    field_alias = f"{lyr_type} Acres"
                    arcpy.management.AddField(new_layer, "Acres", "FLOAT", 2, 2, field_alias=field_alias)

                # Calculate geometry
                arcpy.management.CalculateGeometryAttributes(in_features=new_layer.name, geometry_property=[["Acres", "AREA_GEODESIC"]], area_unit="ACRES_US")

                # Update soils clip layer symbology
                sym = new_layer.symbology
                sym.renderer.symbol.color = {'RGB' : [0, 0, 0, 0]}
                sym.renderer.symbol.outlineColor = {'RGB' : [255, 255, 0, 100]}
                sym.renderer.symbol.size = 1.5
                new_layer.symbology = sym

                # Add label
                new_layer.showLabels = True
                label_class = new_layer.listLabelClasses()[0]
                label_class.visible = True
                label_class.expression = f"$feature.{soils_musym}"

                l_cim = new_layer.getDefinition('V3')
                lc = l_cim.labelClasses[0]

                # Update text properties of label
                lc.textSymbol.symbol.height = 12
                lc.textSymbol.symbol.symbol.symbolLayers = [
                    {
                        "type": "CIMSolidFill",
                        "enable": True,
                        "color": {
                            "type": "CIMRGBColor",
                            "values": [255, 255, 0, 100]
                        }
                    }
                ]
                lc.standardLabelPlacementProperties.numLabelsOption = "OneLabelPerPart"

                # Update CIM definition
                new_layer.setDefinition(l_cim)

                # Get soils layer attribute table and export / extract needed fields for layout
                table_path = "{}\\{}".format(arcpy.env.workspace, f"{sanitize(new_layer_name)}_ExportTable")
                arcpy.conversion.ExportTable(new_layer.name, table_path)
                arcpy.management.DeleteField(table_path, [f"{soils_musym}", "Acres", f"{soils_mukey}"], "KEEP_FIELDS")

                # Add soils table export to the given map
                soils_table = arcpy.mp.Table(table_path)
                tables.append(soils_table)
                m.addTable(soils_table)
                soils_table_uri = soils_table.URI

                # Get layout table
                tbl = lyt.listElements("MAPSURROUND_ELEMENT", lyr_type)[0]

                # Set layout table to exported attributes table
                tbl_cim = tbl.getDefinition("V3")
                tbl_cim.mapMemberURI = soils_table_uri
                tbl.setDefinition(tbl_cim)

                # Refresh layout
                lyt_cim = lyt.getDefinition('V3')
                lyt.setDefinition(lyt_cim)

                # save and close layouts
                project.save()
                project.closeViews("LAYOUTS")

            # Reorder layers so soils layers are last
            log(f"reordering layers for {parcel}")
            for soils_layer in soils_layers:
                for use_layer in use_layers:
                    m.moveLayer(use_layer, soils_layer, "AFTER")

            # Remove unused layout tables
            log(f"removing unused tables for {parcel}")
            uses = {'Agland', 'Forest', 'NonAg'}
            for i in uses:
                if i not in lyr_types:
                    tbls = lyt.listElements("MAPSURROUND_ELEMENT", i)
                    if len(tbls) > 0:
                        tbl_remove = tbls[0]
                        lyt.deleteElement(tbl_remove)

            # Display wanted legend items only
            log(f"removing unused legend items for {parcel}")
            legend = lyt.listElements("LEGEND_ELEMENT")[0]
            legend_items = legend.items
            use_layer_names = [ i.name for i in use_layers ]
            for item in legend_items:
                if item.name in use_layer_names:
                    item.visible = True
                else:
                    item.visible = False

            # Populate soil group worksheet with values from tables
            log(f"filling out {parcel} soil group worksheet")
            sgw_path = f"{output_folder}\\{lyt.name}.xlsx"
            sgw_path = pathlib.PureWindowsPath(sgw_path).as_posix()
            sgw_workbook = openpyxl.load_workbook(sgw_path)
            ws = sgw_workbook['SGW']
            for table in tables:
                table_name = table.name.lower()
                with arcpy.da.SearchCursor(table, ["MUSYM", "MUKEY", "Acres"]) as cursor:
                    idx = 0
                    tot = 0
                    for row in cursor:
                        musym = row[0]
                        mukey = int(row[1])
                        acres = round(float(row[2]), 2)

                        if "agland" in table_name:
                            if idx < 24:
                                soil_cell = f'A{34 + idx}'
                                area_cell = f'H{34 + idx}'
                                mukey_cell = f'F{34 + idx}'
                                ws[soil_cell] = musym
                                ws[mukey_cell] = mukey
                                ws[area_cell] = acres
                            else:
                                # overflow
                                soil_cell = f'N{9 + idx}'
                                area_cell = f'U{9 + idx}'
                                mukey_cell = f'S{9 + idx}'
                                ws[soil_cell] = musym
                                ws[mukey_cell] = mukey
                                ws[area_cell] = acres
                            idx += 1
                        else:
                            tot += acres
                    if "forest" in table_name:
                        ws['L24'] = round(tot, 2)
                    elif "nonag" in table_name:
                        ws['K28'] = round(tot, 2)
            sgw_workbook.save(sgw_path)
            sgw_workbook.close()
            del sgw_workbook
            del ws

        # open layouts
        log("opening layouts")
        for layout in layouts:
            layout.openView()

        # Save
        log("saving project")
        project.save()
        del project
